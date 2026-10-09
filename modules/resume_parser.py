from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path


class DocumentParseError(ValueError):
    """A supported document could not be read or contained no usable text."""


@dataclass(frozen=True)
class ParsedDocument:
    text: str
    file_name: str
    file_type: str
    page_count: int | None = None


def normalize_text(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"(?<=\w)-\s*\n\s*(?=\w)", "", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def parse_document(
    file_name: str,
    file_bytes: bytes,
    max_bytes: int = 10 * 1024 * 1024,
    enable_ocr: bool = False,
) -> ParsedDocument:
    if not file_bytes:
        raise DocumentParseError("The uploaded file is empty.")
    if len(file_bytes) > max_bytes:
        raise DocumentParseError(
            f"The file exceeds the {max_bytes // (1024 * 1024)} MB upload limit."
        )

    extension = Path(file_name).suffix.lower()
    if extension == ".pdf" and b"%PDF-" not in file_bytes[:1024]:
        raise DocumentParseError("The uploaded file does not have a valid PDF signature.")
    if extension == ".docx":
        if not file_bytes.startswith(b"PK\x03\x04"):
            raise DocumentParseError("The uploaded file does not have a valid DOCX signature.")
        try:
            with zipfile.ZipFile(io.BytesIO(file_bytes)) as archive:
                members = archive.infolist()
                total_uncompressed = sum(item.file_size for item in members)
                if (
                    len(members) > 2000
                    or total_uncompressed > 50 * 1024 * 1024
                    or "[Content_Types].xml" not in archive.namelist()
                    or "word/document.xml" not in archive.namelist()
                ):
                    raise DocumentParseError(
                        "The DOCX package is invalid or exceeds safe extraction limits."
                    )
        except zipfile.BadZipFile as error:
            raise DocumentParseError("The uploaded DOCX package is damaged.") from error
    try:
        if extension == ".pdf":
            import fitz

            with fitz.open(stream=file_bytes, filetype="pdf") as pdf:
                text = "\n\n".join(page.get_text("text") for page in pdf)
                page_count = len(pdf)
                if enable_ocr and not text.strip():
                    try:
                        text = "\n\n".join(
                            page.get_text(
                                "text",
                                textpage=page.get_textpage_ocr(language="eng", dpi=200),
                            )
                            for page in pdf
                        )
                    except Exception as error:
                        raise DocumentParseError(
                            "PDF OCR failed. Install Tesseract OCR and its English "
                            "language data, then retry; selectable PDF text extraction "
                            "continues to work without OCR."
                        ) from error
            file_type = "PDF"
        elif extension == ".docx":
            from docx import Document

            document = Document(io.BytesIO(file_bytes))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                paragraphs.extend(
                    " | ".join(cell.text for cell in row.cells) for row in table.rows
                )
            text = "\n".join(paragraphs)
            page_count = None
            file_type = "DOCX"
        elif extension in {".txt", ".md", ".rtf"}:
            page_count = None
            file_type = "TEXT"
            for encoding_name in ("utf-8", "utf-16", "cp1252", "latin-1"):
                try:
                    text = file_bytes.decode(encoding_name)
                    break
                except UnicodeDecodeError:
                    continue
            else:
                text = file_bytes.decode("utf-8", errors="replace")
        else:
            raise DocumentParseError(
                "Unsupported file type. Upload a PDF, DOCX, or text document."
            )
    except DocumentParseError:
        raise
    except Exception as error:
        raise DocumentParseError(
            f"Could not read this {extension[1:].upper()} file. It may be damaged "
            "or password-protected."
        ) from error

    cleaned = normalize_text(text)
    if len(cleaned) > 1_000_000:
        raise DocumentParseError(
            "The extracted document text exceeds the 1,000,000-character safety limit."
        )
    if not cleaned:
        message = "No selectable text was found in this PDF. It may be a scanned image; OCR is required."
        if extension == ".docx":
            message = "No readable text was found in this DOCX file."
        elif extension in {".txt", ".md", ".rtf"}:
            message = "No readable text was found in this text document."
        raise DocumentParseError(message)

    return ParsedDocument(cleaned, Path(file_name).name, file_type, page_count)
